import os
from docx import Document
from config import WORD_TEMPLATE_PATH

def generate_word_documents(rules, output_folder):
    output_files = []
    try:
        for rule in rules:
            # Load the Word template
            doc = Document(WORD_TEMPLATE_PATH)
            
            # Replace placeholders with rule information
            for paragraph in doc.paragraphs:
                replace_placeholder(paragraph, '{RuleID}', rule['rule_id'])
                replace_placeholder(paragraph, '{Severity}', rule['severity'])
                replace_placeholder(paragraph, '{Description}', rule['description'])
            
            # Save the document
            output_filename = f"Rule_{rule['rule_id']}.docx"
            output_path = os.path.join(output_folder, output_filename)
            doc.save(output_path)
            output_files.append(output_filename)
        
        return output_files
    except Exception as e:
        raise Exception(f"Error generating Word documents: {str(e)}")

def replace_placeholder(paragraph, placeholder, value):
    if placeholder in paragraph.text:
        for run in paragraph.runs:
            run.text = run.text.replace(placeholder, str(value))
